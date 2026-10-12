"""
📥 قراءة فاتورة شراء من ملف: Excel (.xlsx/.xls) أو CSV أو PDF أو صورة.

كل الصيغ بتطلع نفس الشكل: ``(اسم المورد, [بنود])`` — كل بند
``{part_number, name, qty, cost, sell}`` — وبعدها صفحة الاستيراد بتطابقها
بالمنتجات وتعرضها للمراجعة قبل الحفظ.

* **الجدول**: صف العناوين مش لازم يكون أول صف (فواتير الموردين وملفات
  التصدير فيها عنوان وتاريخ فوق الجدول) — بندوّر عليه في أول ٢٥ صف.
  العناوين بالعربي أو الإنجليزي (كود/رقم القطعة، البيان/اسم الصنف، العدد/الكمية،
  السعر/الفئة/سعر الشراء، الإجمالي/القيمة، سعر البيع). صفوف الإجمالي والخصم
  والضريبة بتتشال، ولو مفيش سعر وحدة بنقسم الإجمالي على الكمية.
* **PDF من برنامج** (فيه نص): بنقرا الجدول من النص مباشرة من غير ذكاء اصطناعي.
  النص العربي في الـ PDF بيطلع معكوس (ترتيب العرض) فبنرجّعه للترتيب الصحيح.
* **PDF ممسوح / صورة**: بيروح لقارئ الفواتير بالذكاء الاصطناعي لو مفعّل.
"""
from __future__ import annotations

import base64
import io
import re
import unicodedata
from decimal import Decimal, InvalidOperation

from django.conf import settings

MAX_UPLOAD_BYTES = 15 * 1024 * 1024
IMAGE_EXTS = ('.png', '.jpg', '.jpeg', '.webp', '.gif', '.bmp', '.heic', '.heif')
HEADER_SCAN_ROWS = 25


class ImportFileError(Exception):
    """رسالة واضحة للمستخدم (بتترجّع 400 من صفحة الاستيراد)."""


# ------------------------------------------------------------------ نصوص وأرقام
_TASHKEEL = re.compile(r'[ؐ-ًؚ-ٰٟۖ-ۭـ]')
_AR_LETTER = re.compile(r'[ء-يٱ-ۓﭐ-﷿ﹰ-ﻼ]')
_DIGITS = str.maketrans('٠١٢٣٤٥٦٧٨٩۰۱۲۳۴۵۶۷۸۹٫٬', '01234567890123456789.,')
_NUM = re.compile(r'-?\d[\d,.]*')


def _norm(text) -> str:
    """تطبيع للمقارنة: من غير تشكيل، أ/إ/آ → ا، ى → ي، ة → ه، حروف صغيرة."""
    s = unicodedata.normalize('NFKC', str(text or ''))
    s = _TASHKEEL.sub('', s)
    s = s.translate(str.maketrans('أإآٱىة', 'اااايه'))
    return ' '.join(s.lower().split())


def _cell_text(v) -> str:
    if v is None:
        return ''
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return ' '.join(str(v).split())


def parse_number(v):
    """«1,250.50» / «14300,00» / «١٢٥٠» / «850 ج.م» / 4.0 → Decimal (أو None)."""
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, (int, float, Decimal)):
        try:
            return Decimal(str(v))
        except InvalidOperation:
            return None
    m = _NUM.search(str(v).translate(_DIGITS))
    if not m:
        return None
    s = m.group(0).rstrip('.,')
    if ',' in s and '.' in s:
        # آخر فاصل هو العلامة العشرية: 1,250.50 أو 1.250,50
        s = s.replace(',', '') if s.rfind('.') > s.rfind(',') else s.replace('.', '').replace(',', '.')
    elif ',' in s:
        head, _sep, tail = s.rpartition(',')
        s = f"{head.replace(',', '')}.{tail}" if (s.count(',') == 1 and len(tail) in (1, 2)) else s.replace(',', '')
    elif s.count('.') > 1:
        s = s.replace('.', '')
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def _fix_pdf_text(s: str) -> str:
    """النص العربي اللي بيطلع من PDF مترتّب زي ما بيتعرض (شمال → يمين) وبحروف
    الأشكال المتصلة — بنرجّعه للترتيب الطبيعي (bidi) وبعدين للحروف العادية."""
    s = _cell_text(s)
    if not s or not _AR_LETTER.search(s):
        return s
    try:
        from bidi import get_display
        s = get_display(s, base_dir='R')
    except Exception:  # noqa: BLE001 — مكتبة اختيارية؛ الأسوأ نص معكوس قابل للتعديل
        s = s[::-1]
    return unicodedata.normalize('NFKC', s)


_CODE_SEP = re.compile(r'([\-\s/._]+)')


def flipped_code(code: str) -> str:
    """كود القطعة في PDF اتعمل من صفحة عربي (RTL) ممكن يطلع أجزاؤه مقلوبة:
    90915-YZZE1 بيتعرض YZZE1-90915. الشكل التاني بنجرّبه في المطابقة بس."""
    parts = _CODE_SEP.split(code or '')
    return ''.join(reversed(parts)) if len(parts) > 1 else code


# ------------------------------------------------------------------ عناوين الأعمدة
# بالترتيب: أول مجموعة تنطبق على العنوان هي اللي بتاخده.
_FIELD_KEYS = (
    # أعمدة بنتجاهلها (حد التنبيه، توزيع الفروع في تصدير المخزون)
    ('skip', ('تنبيه', 'reorder', 'min stock', 'الفروع', 'branches')),
    ('total', ('اجمالي', 'الاجمالي', 'total', 'amount', 'القيمه', 'قيمه', 'المبلغ', 'subtotal', 'value')),
    ('wholesale', ('جمله', 'wholesale', 'b2b')),
    ('sell', ('بيع', 'sale', 'selling', 'retail', 'قطاعي')),
    ('sku', ('كود', 'code', 'part', 'sku', 'قطعه', 'رقم الصنف', 'barcode', 'باركود', 'oem', 'item no', 'ref')),
    # «المخزون» = عمود الكمية في تصدير المخزون (نقل بضاعة من فرع/شركة لتانية)
    ('qty', ('كميه', 'عدد', 'qty', 'quantity', 'pcs', 'quant', 'مخزون', 'رصيد', 'stock', 'on hand', 'متاح')),
    ('cost', ('شراء', 'تكلفه', 'cost', 'سعر', 'price', 'فئه', 'rate')),
    ('name', ('اسم', 'صنف', 'بيان', 'وصف', 'name', 'desc', 'item', 'product', 'منتج', 'بند')),
)
_FIELD_KEYS = tuple((f, tuple(_norm(k) for k in keys)) for f, keys in _FIELD_KEYS)
_WHOLESALE_KEYS = dict(_FIELD_KEYS)['wholesale']
_SELL_KEYS = dict(_FIELD_KEYS)['sell']


def _classify(h):
    """الحقل اللي العنوان ده بيمثّله (أو None)."""
    for field, keys in _FIELD_KEYS:
        if any(k in h for k in keys):
            if field == 'total':
                # «قيمة المخزون (جملة)» / «(بيع)» = إجمالي بسعر الجملة/البيع مش بالتكلفة
                if any(k in h for k in _WHOLESALE_KEYS):
                    return 'wholesale_total'
                if any(k in h for k in _SELL_KEYS):
                    return 'sell_total'
            return None if field == 'skip' else field
    return None


# صف إجمالي/خصم/ضريبة تحت الجدول (كلمة كاملة في أول الخانة — «Taxi» مش ضريبة)
_SUMMARY_ROW = re.compile(r'(?:%s)(?:\W|$)' % '|'.join(re.escape(_norm(w)) for w in (
    'اجمالي', 'الاجمالي', 'المجموع', 'الصافي', 'المطلوب', 'total', 'subtotal', 'grand total',
    'خصم', 'الخصم', 'discount', 'ضريبه', 'الضريبه', 'tax', 'vat')))


def _header_fields(row):
    """{field: column index} لصف مرشّح يبقى صف العناوين."""
    found = {}
    for idx, cell in enumerate(row):
        h = _norm(cell)
        if not h or len(h) > 40:
            continue
        field = _classify(h)
        if field:
            found.setdefault(field, idx)
    return found


_NUMERIC = {'qty', 'cost', 'total', 'sell', 'wholesale', 'sell_total', 'wholesale_total'}


def _find_header(rows):
    best, best_score = None, 0
    for i, row in enumerate(rows[:HEADER_SCAN_ROWS]):
        f = _header_fields(row)
        if not ({'sku', 'name'} & f.keys()):
            continue
        score = len(f) + (2 if (_NUMERIC & f.keys()) else 0)
        if score > best_score:
            best, best_score = (i, f), score
    return best


def rows_to_items(rows):
    """صفوف جدول (أي صيغة) → بنود فاتورة شراء."""
    rows = [list(r or []) for r in rows if r is not None]
    found = _find_header(rows)
    if found is None:
        return []
    start, cols = found
    # مفيش سعر شراء ولا إجمالي بالتكلفة: سعر الجملة (نقل بضاعة بين شركتين/فرعين) أو
    # سعر البيع (فاتورة مورد بـ«سعر البيع» بتاعه) بيبقى سعر الشراء — ومتعلّم للمراجعة.
    no_cost = not ({'cost', 'total'} & cols.keys())
    has_numbers = bool(_NUMERIC & cols.keys())
    header_sig = [_norm(c) for c in rows[start]]

    def cell(row, field):
        idx = cols.get(field)
        return row[idx] if idx is not None and idx < len(row) else None

    def per_unit(price, line_total, qty):
        if price and price > 0:
            return price
        if line_total and line_total > 0:
            return (line_total / qty).quantize(Decimal('0.01'))
        return None

    items = []
    for row in rows[start + 1:]:
        if [_norm(c) for c in row] == header_sig:   # عناوين متكررة في كل صفحة PDF
            continue
        sku = _cell_text(cell(row, 'sku'))
        name = _cell_text(cell(row, 'name'))
        if not sku and not name:
            continue
        if _SUMMARY_ROW.match(_norm(name or sku)):
            continue
        nums = {f: parse_number(cell(row, f)) for f in _NUMERIC}
        if has_numbers and all(v is None for v in nums.values()):
            continue   # عنوان قسم / ملاحظة جوه الجدول
        qty = nums['qty']
        if qty is not None and qty < 1:
            continue   # صنف رصيده صفر (تصدير مخزون) — مفيش حاجة تتشترى
        qty = int(qty) if qty is not None else 1
        cost = per_unit(nums['cost'], nums['total'], qty)
        sell = per_unit(nums['sell'], nums['sell_total'], qty)
        whole = per_unit(nums['wholesale'], nums['wholesale_total'], qty)
        guess = ''
        if no_cost and whole:
            cost, guess = whole, 'wholesale'
        elif no_cost and sell:
            cost, guess = sell, 'sell'
        else:
            # سعر المورد نفسه في عمود «سعر البيع/الجملة» بتاعه = الشراء بتاعنا، مش أسعارنا
            if sell and cost and abs(sell - cost) < Decimal('0.01'):
                sell = None
            if whole and cost and abs(whole - cost) < Decimal('0.01'):
                whole = None
        item = {
            'part_number': sku[:100], 'name': name[:255], 'qty': qty,
            'cost': float(cost) if cost and cost > 0 else 0,
            'sell': float(sell) if sell and sell > 0 else 0,
            'wholesale': float(whole) if whole and whole > 0 else 0,
        }
        if guess:
            item['cost_guess'] = guess
        items.append(item)
    return items


# ------------------------------------------------------------------ قراءة الملفات
def _xlsx_sheets(up):
    import openpyxl
    try:
        wb = openpyxl.load_workbook(up, read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001
        raise ImportFileError('ملف Excel ده مش سليم أو محمي بباسورد — افتحه واحفظه تاني بصيغة .xlsx.') from exc
    active = wb.active
    sheets = [active] + [ws for ws in wb.worksheets if ws is not active]
    return [list(ws.iter_rows(values_only=True)) for ws in sheets]


def _xls_sheets(up):
    try:
        import xlrd
        book = xlrd.open_workbook(file_contents=up.read())
    except Exception as exc:  # noqa: BLE001
        raise ImportFileError('ملف Excel القديم (.xls) ده مش مقروء — احفظه بصيغة .xlsx وارفعه تاني.') from exc
    return [[sh.row_values(i) for i in range(sh.nrows)] for sh in book.sheets()]


def _csv_rows(up):
    import csv
    raw = up.read()
    for enc in ('utf-8-sig', 'cp1256'):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    else:
        text = raw.decode('utf-8', errors='ignore')
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=',;\t|')
    except csv.Error:
        dialect = csv.excel
    return list(csv.reader(io.StringIO(text), dialect))


def _pdf_rows(data: bytes):
    """صفوف الجداول من PDF فيه نص (مش ممسوح). [] لو مفيش نص/جداول."""
    try:
        import pdfplumber
    except ImportError:
        return []
    import logging
    logging.getLogger('pdfminer').setLevel(logging.ERROR)   # تحذيرات خطوط مالهاش لازمة في اللوج
    rows = []
    try:
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            for page in pdf.pages[:30]:
                tables = page.extract_tables()
                if not tables:   # جدول من غير خطوط → نقسّم بالمسافات
                    tables = page.extract_tables({'vertical_strategy': 'text', 'horizontal_strategy': 'text'})
                for t in tables:
                    rows.extend([_fix_pdf_text(c) for c in r] for r in t)
    except Exception:  # noqa: BLE001 — PDF تالف/مشفّر
        return []
    return rows


def vision_ai_ready() -> bool:
    return bool(getattr(settings, 'ENABLE_AI_PREDICTIONS', False)
                and (str(getattr(settings, 'AI_VISION_API_KEY', '') or '').strip()
                     or str(getattr(settings, 'GEMINI_API_KEY', '') or '').strip()))


def _ai_items(data: bytes, mime: str):
    from inventory.ai_services import scan_invoice_image_ai
    try:
        res = scan_invoice_image_ai(base64.b64encode(data).decode(), mime_type=mime) or {}
    except Exception:  # noqa: BLE001
        res = {}
    items = []
    for it in res.get('items') or []:
        qty = parse_number(it.get('qty'))
        cost = parse_number(it.get('cost'))
        items.append({
            'part_number': _cell_text(it.get('part_number'))[:100],
            'name': _cell_text(it.get('name'))[:255],
            'qty': int(qty) if qty and qty >= 1 else 1,
            'cost': float(cost) if cost and cost > 0 else 0,
            'sell': 0, 'wholesale': 0,
        })
    vendor = _cell_text(res.get('vendor_name'))
    return ('' if vendor == 'مجهول' else vendor), items


_NO_AI = ('قراءة {what} بالذكاء الاصطناعي مش مفعّلة على السيستم — '
          'ارفع ملف Excel أو PDF طالع من برنامج المورد، أو نزّل النموذج واملاه.')
_NO_TABLE = ('مالقيناش جدول أصناف في الملف — اتأكد إن فيه صف عناوين زي '
             '«الكود / الصنف / الكمية / السعر» (أو نزّل النموذج واملاه).')


def extract_purchase_items(up):
    """(اسم المورد, [بنود]) من الملف المرفوع — أو ImportFileError برسالة للمستخدم."""
    name = (getattr(up, 'name', '') or '').lower()
    if (getattr(up, 'size', 0) or 0) > MAX_UPLOAD_BYTES:
        raise ImportFileError('الملف أكبر من 15 ميجا — قسّمه أو صغّر الصورة.')

    if name.endswith(IMAGE_EXTS):
        if not vision_ai_ready():
            raise ImportFileError(_NO_AI.format(what='الصور'))
        vendor, items = _ai_items(up.read(), 'image/jpeg')
        if not items:
            raise ImportFileError('مقدرناش نقرا بنود من الصورة — صوّرها أوضح (الفاتورة كاملة ومفرودة) أو ارفع Excel/PDF.')
        return vendor, items

    if name.endswith('.pdf'):
        data = up.read()
        items = rows_to_items(_pdf_rows(data))
        if items:
            return '', items
        if not vision_ai_ready():
            raise ImportFileError(
                'الـ PDF ده مفيهوش جدول نصّي نقدر نقراه (غالباً صورة ممسوحة). ' + _NO_AI.format(what='الملفات الممسوحة'))
        vendor, items = _ai_items(data, 'application/pdf')
        if not items:
            raise ImportFileError('مقدرناش نقرا بنود من الـ PDF — جرّب صورة أوضح أو ملف Excel.')
        return vendor, items

    if name.endswith(('.csv', '.txt')):
        sheets = [_csv_rows(up)]
    elif name.endswith('.xls'):
        sheets = _xls_sheets(up)
    elif name.endswith(('.xlsx', '.xlsm')):
        sheets = _xlsx_sheets(up)
    else:
        raise ImportFileError('نوع الملف ده مش مدعوم — ارفع Excel أو CSV أو PDF أو صورة.')
    for rows in sheets:
        items = rows_to_items(rows)
        if items:
            return '', items
    raise ImportFileError(_NO_TABLE)


def template_xlsx() -> bytes:
    """نموذج Excel فاضي بالعناوين اللي القارئ بيفهمها + مثالين."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = 'فاتورة شراء'
    ws.sheet_view.rightToLeft = True
    head = ['رقم القطعة', 'اسم الصنف', 'الكمية', 'سعر الشراء', 'سعر البيع (اختياري)']
    ws.append(head)
    for c in ws[1]:
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = PatternFill('solid', fgColor='EA580C')
    ws.append(['04465-0D130', 'تيل فرامل أمامي', 4, 850, 1100])
    ws.append(['90915-YZZE1', 'فلتر زيت', 10, 125.5, None])
    for col, width in zip('ABCDE', (22, 36, 12, 14, 20)):
        ws.column_dimensions[col].width = width
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
