# 📤 تصدير التقارير إلى Excel (.xlsx) و PDF — مُوحّد لكل التقارير.
#
# كل تقرير بيبني بنية موحّدة `tables` = قائمة جداول:
#   [{"name": "...", "columns": [...], "rows": [[...], ...], "total": [...]}]
# وبعدين بننادي report_to_xlsx / report_to_pdf.
from decimal import Decimal

from django.http import HttpResponse


def _cell_value(v):
    """أرقام Decimal بتتحوّل float عشان Excel يتعامل معها كأرقام.

    🐛 [FIX]: نصوص الترجمة (get_FOO_display / gettext_lazy) كانت بتوقّع تصدير
    Excel بـ 500 («Cannot convert 'بيع قطع غيار' to Excel») — بتتحوّل نص عادي.
    والتواريخ بمنطقة زمنية (Excel مابيقبلهاش) بتتحوّل لتوقيت المحل.
    """
    from datetime import datetime
    from django.utils import timezone
    from django.utils.functional import Promise
    if isinstance(v, Decimal):
        return float(v)
    if isinstance(v, Promise):
        return str(v)
    if isinstance(v, datetime) and timezone.is_aware(v):
        return timezone.localtime(v).replace(tzinfo=None)
    return v


def report_to_xlsx(filename, title, subtitle, tables):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "التقرير"
    ws.sheet_view.rightToLeft = True

    r = 1
    ws.cell(r, 1, _cell_value(title)).font = Font(bold=True, size=14)
    r += 1
    if subtitle:
        ws.cell(r, 1, _cell_value(subtitle)).font = Font(size=10, color="666666")
        r += 1
    r += 1

    max_cols = 1
    hdr_fill = PatternFill("solid", fgColor="1F2937")
    for t in tables:
        max_cols = max(max_cols, len(t.get("columns") or []))
        if t.get("name"):
            ws.cell(r, 1, _cell_value(t["name"])).font = Font(bold=True, size=12)
            r += 1
        for c, col in enumerate(t.get("columns") or [], 1):
            cell = ws.cell(r, c, _cell_value(col))
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = hdr_fill
        r += 1
        for row in t.get("rows") or []:
            for c, val in enumerate(row, 1):
                ws.cell(r, c, _cell_value(val))
            r += 1
        if t.get("total"):
            for c, val in enumerate(t["total"], 1):
                cell = ws.cell(r, c, _cell_value(val))
                cell.font = Font(bold=True)
            r += 1
        r += 1

    for c in range(1, max_cols + 1):
        ws.column_dimensions[get_column_letter(c)].width = 24

    resp = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    resp["Content-Disposition"] = f'attachment; filename="{filename}.xlsx"'
    wb.save(resp)
    return resp


def _reportlab_pdf(title, subtitle, tables):
    """PDF جدولي من اليمين للشمال بخط عربي (Cairo/Amiri) — من غير weasyprint."""
    import io
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    from erp_core.ai.print_spec_pdf import _resolve_arabic_font_paths, _shape_arabic_for_pdf

    import os
    from django.conf import settings
    # Amiri فيه كل أشكال الحروف بعد الـ reshape (Cairo المتغيّر ناقصه «لا» و«أ»…)
    amiri = os.path.join(settings.BASE_DIR, 'static', 'fonts', 'Amiri-Bold.ttf')
    reg, bold = _resolve_arabic_font_paths()
    if os.path.exists(amiri):
        reg = bold = amiri
    font, font_bold = 'Helvetica', 'Helvetica-Bold'
    if reg:
        try:
            pdfmetrics.registerFont(TTFont('RepAr', reg))
            pdfmetrics.registerFont(TTFont('RepArB', bold or reg))
            font, font_bold = 'RepAr', 'RepArB'
        except Exception:
            pass

    def txt(v):
        v = _cell_value(v)
        if v is None:
            return ''
        if isinstance(v, float):
            return f"{v:,.2f}"
        if hasattr(v, 'strftime'):
            return v.strftime('%Y-%m-%d %H:%M')
        return _shape_arabic_for_pdf(str(v))

    max_cols = max([len(t.get("columns") or []) for t in tables] or [1])
    page = landscape(A4) if max_cols > 6 else A4
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=page, leftMargin=12 * mm, rightMargin=12 * mm,
                            topMargin=12 * mm, bottomMargin=12 * mm, title=str(title))
    h1 = ParagraphStyle('h1', fontName=font_bold, fontSize=16, alignment=2, leading=22)
    sub = ParagraphStyle('sub', fontName=font, fontSize=9, alignment=2, textColor=colors.grey, leading=13)
    tname = ParagraphStyle('tn', fontName=font_bold, fontSize=11, alignment=2, leading=16)
    story = [Paragraph(txt(title), h1)]
    if subtitle:
        story.append(Paragraph(txt(subtitle), sub))
    story.append(Spacer(1, 6 * mm))
    for t in tables:
        if t.get("name"):
            story.append(Paragraph(txt(t["name"]), tname))
            story.append(Spacer(1, 2 * mm))
        cols = list(t.get("columns") or [])
        # RTL: أول عمود على اليمين
        data = [[txt(c) for c in reversed(cols)]]
        for row in t.get("rows") or []:
            data.append([txt(v) for v in reversed(list(row))])
        if t.get("total"):
            data.append([txt(v) for v in reversed(list(t["total"]))])
        if not cols:
            continue
        tbl = Table(data, repeatRows=1, hAlign='RIGHT')
        style = [
            ('FONTNAME', (0, 0), (-1, -1), font), ('FONTSIZE', (0, 0), (-1, -1), 8),
            ('FONTNAME', (0, 0), (-1, 0), font_bold),
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1f2937')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('ALIGN', (0, 0), (-1, -1), 'RIGHT'), ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('GRID', (0, 0), (-1, -1), 0.25, colors.HexColor('#d1d5db')),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f9fafb')]),
        ]
        if t.get("total"):
            style += [('FONTNAME', (0, -1), (-1, -1), font_bold),
                      ('BACKGROUND', (0, -1), (-1, -1), colors.HexColor('#e5e7eb'))]
        tbl.setStyle(TableStyle(style))
        story += [tbl, Spacer(1, 6 * mm)]
    doc.build(story)
    return buf.getvalue()


def report_to_pdf(request, filename, title, subtitle, tables):
    from django.template.loader import render_to_string
    html = render_to_string("inventory/report_print.html", {
        "title": title, "subtitle": subtitle, "tables": tables,
    })
    try:
        from weasyprint import HTML
        pdf = HTML(string=html, base_url=request.build_absolute_uri('/')).write_pdf()
    except Exception:
        # 🐛 [FIX]: weasyprint مش في requirements أصلاً — فالـ PDF عمره ما كان
        #    بيطلع على السيرفر (بيرجع صفحة HTML). reportlab موجود: نعمل بيه
        #    PDF حقيقي بالعربي، ولو فشل لأي سبب نرجّع الـ HTML للطباعة.
        try:
            pdf = _reportlab_pdf(title, subtitle, tables)
        except Exception:
            import logging
            logging.getLogger('mouss_tec_core').exception("[PDF] reportlab fallback failed")
            return HttpResponse(html)
    resp = HttpResponse(pdf, content_type="application/pdf")
    resp["Content-Disposition"] = f'attachment; filename="{filename}.pdf"'
    return resp


def export_report(request, filename, title, subtitle, tables):
    """يرجّع Response حسب ?export=xlsx|pdf، أو None لو مفيش تصدير مطلوب."""
    exp = (request.GET.get("export") or "").strip().lower()
    if exp == "xlsx":
        return report_to_xlsx(filename, title, subtitle, tables)
    if exp == "pdf":
        return report_to_pdf(request, filename, title, subtitle, tables)
    return None
